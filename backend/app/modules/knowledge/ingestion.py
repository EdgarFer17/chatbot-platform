import hashlib
import ipaddress
import io
import re
import socket
import zipfile
import xml.etree.ElementTree as ET
from dataclasses import dataclass
from pathlib import PurePosixPath
from urllib.parse import unquote, urljoin, urlparse

import httpx
from bs4 import BeautifulSoup
from docx import Document as DocxDocument
from pypdf import PdfReader

MAX_DOWNLOAD_BYTES = 20 * 1024 * 1024
MAX_DISCOVERED_DOCUMENTS = 40
MAX_WEBSITE_PAGES = 500
SUPPORTED_EXTENSIONS = {"", ".pdf", ".docx", ".txt", ".md", ".mdx", ".htm", ".html"}


class IngestionError(ValueError):
    pass


@dataclass(frozen=True)
class ExtractedDocument:
    title: str
    content: str
    url: str
    mime_type: str
    checksum: str


def validate_public_url(raw_url: str, expected_host: str | None = None) -> tuple[str, str]:
    parsed = urlparse(raw_url.strip())
    if parsed.scheme not in {"http", "https"} or not parsed.hostname or parsed.username or parsed.password:
        raise IngestionError("Use uma URL pública HTTP ou HTTPS")
    host = parsed.hostname.lower().rstrip(".")
    if expected_host and host != expected_host:
        raise IngestionError("A sincronização não pode sair do domínio autorizado")
    try:
        addresses = {item[4][0] for item in socket.getaddrinfo(host, parsed.port or (443 if parsed.scheme == "https" else 80), type=socket.SOCK_STREAM)}
    except socket.gaierror as error:
        raise IngestionError("Não foi possível resolver o domínio informado") from error
    for address in addresses:
        ip = ipaddress.ip_address(address)
        if ip.is_private or ip.is_loopback or ip.is_link_local or ip.is_reserved or ip.is_unspecified or ip.is_multicast:
            raise IngestionError("O endereço informado não é uma fonte pública permitida")
    return raw_url.strip(), host


def fetch_public_document(url: str, allowed_host: str) -> tuple[bytes, str, str]:
    current = url
    with httpx.Client(timeout=25, headers={"User-Agent": "ChatbotPlatformKnowledge/0.1"}) as client:
        for _ in range(4):
            current, _ = validate_public_url(current, allowed_host)
            with client.stream("GET", current, follow_redirects=False) as response:
                if response.status_code in {301, 302, 303, 307, 308}:
                    location = response.headers.get("location")
                    if not location:
                        raise IngestionError("Redirecionamento inválido")
                    current = urljoin(current, location)
                    continue
                response.raise_for_status()
                content_length = int(response.headers.get("content-length", "0") or 0)
                if content_length > MAX_DOWNLOAD_BYTES:
                    raise IngestionError("O documento excede o limite de 20 MB")
                data = bytearray()
                for chunk in response.iter_bytes():
                    data.extend(chunk)
                    if len(data) > MAX_DOWNLOAD_BYTES:
                        raise IngestionError("O documento excede o limite de 20 MB")
                return bytes(data), response.headers.get("content-type", "application/octet-stream").split(";", 1)[0].lower(), str(response.url)
    raise IngestionError("A fonte excedeu o limite de redirecionamentos")


def extract_document(data: bytes, mime_type: str, url: str, title_hint: str | None = None) -> ExtractedDocument:
    if not data or len(data) > MAX_DOWNLOAD_BYTES:
        raise IngestionError("O arquivo está vazio ou excede o limite de 20 MB")
    try:
        return _extract_document(data, mime_type, url, title_hint)
    except IngestionError:
        raise
    except Exception as error:
        raise IngestionError("Não foi possível ler o arquivo. Verifique se está íntegro, sem senha e contém texto selecionável.") from error


def _extract_document(data: bytes, mime_type: str, url: str, title_hint: str | None = None) -> ExtractedDocument:
    extension = PurePosixPath(urlparse(url).path).suffix.lower()
    title = title_hint or unquote(PurePosixPath(urlparse(url).path).name) or "Documento sem título"
    if mime_type == "application/pdf" or extension == ".pdf":
        reader = PdfReader(io.BytesIO(data))
        content = "\n\n".join((page.extract_text() or "").strip() for page in reader.pages).strip()
    elif mime_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document" or extension == ".docx":
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 80 * 1024 * 1024:
                raise IngestionError("O documento descompactado excede o limite permitido")
        document = DocxDocument(io.BytesIO(data))
        content = "\n".join(paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip())
        content += "\n" + "\n".join(" | ".join(cell.text.strip() for cell in row.cells) for table in document.tables for row in table.rows)
    elif mime_type in {"text/html", "application/xhtml+xml"} or extension in {".html", ".htm", ""}:
        soup = BeautifulSoup(data, "html.parser")
        for tag in soup(["script", "style", "noscript", "svg", "nav", "footer", "aside"]):
            tag.decompose()
        for selector in ('[class*="pagination"]', '[class*="breadcrumb"]', '[aria-label*="pagination" i]', '[rel="prev"]', '[rel="next"]'):
            for tag in soup.select(selector):
                tag.decompose()
        content_root = soup.find("main") or soup.find("article") or soup.body or soup
        heading = content_root.find("h1") or soup.find("title")
        if heading:
            title = heading.get_text(" ", strip=True)
        content = re.sub(r"\n{3,}", "\n\n", content_root.get_text("\n", strip=True))
        substantive_content = content
        if title and substantive_content.lower().startswith(title.lower()):
            substantive_content = substantive_content[len(title):].strip()
        if len(substantive_content) < 40:
            raise IngestionError("A página não possui texto principal suficiente; o conteúdo pode estar apenas em imagens")
    elif mime_type.startswith("text/") or extension in {".txt", ".md", ".mdx"}:
        content = data.decode("utf-8", errors="replace").strip()
    else:
        raise IngestionError(f"Formato não suportado: {mime_type}")
    if len(content) < 20:
        raise IngestionError("O documento não possui texto suficiente para indexação")
    return ExtractedDocument(title=title[:220], content=content, url=url, mime_type=mime_type, checksum=hashlib.sha256(content.encode("utf-8")).hexdigest())


def discover_document_urls(html: bytes, base_url: str, allowed_host: str) -> list[tuple[str, str]]:
    soup = BeautifulSoup(html, "html.parser")
    content_root = soup.find("main") or soup.find("article") or soup
    discovered: list[tuple[str, str]] = []
    seen: set[str] = set()
    for link in content_root.find_all("a", href=True):
        href = link["href"].strip()
        if not href or href.startswith(("mailto:", "tel:", "javascript:")):
            continue
        url = urljoin(base_url, href).split("#", 1)[0]
        parsed = urlparse(url)
        extension = PurePosixPath(parsed.path).suffix.lower()
        normalized_url = parsed._replace(query="", fragment="").geturl().rstrip("/") or parsed._replace(query="", fragment="").geturl()
        if parsed.hostname and parsed.hostname.lower().rstrip(".") == allowed_host and extension in SUPPORTED_EXTENSIONS and normalized_url not in seen:
            seen.add(normalized_url)
            discovered.append((normalized_url, link.get_text(" ", strip=True)[:220] or None))
            if len(discovered) >= MAX_DISCOVERED_DOCUMENTS:
                break
    return discovered


def discover_sitemap_urls(xml_data: bytes, allowed_host: str) -> list[str]:
    try:
        root = ET.fromstring(xml_data)
    except ET.ParseError as error:
        raise IngestionError("O sitemap informado não possui XML válido") from error
    urls: list[str] = []
    seen: set[str] = set()
    for element in root.iter():
        if not element.tag.endswith("loc") or not element.text:
            continue
        raw_url = element.text.strip()
        parsed = urlparse(raw_url)
        if not parsed.hostname or parsed.hostname.lower().rstrip(".") != allowed_host:
            continue
        normalized = parsed._replace(query="", fragment="").geturl().rstrip("/") or raw_url
        if normalized not in seen:
            seen.add(normalized)
            urls.append(normalized)
            if len(urls) >= MAX_WEBSITE_PAGES:
                break
    return urls
