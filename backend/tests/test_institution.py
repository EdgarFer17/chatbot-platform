import io
import uuid

import pytest
from docx import Document
from sqlalchemy import func, select

from app.core.config import settings
from app.modules.knowledge.ingestion import extract_document
from app.modules.platform.institution import setup_institution
from app.modules.platform.models import Assistant, ApiKey, Environment, EnvironmentKind, InstitutionConfig, KnowledgeDocument, KnowledgeSource, KnowledgeStatus, SourceType, Tenant

HEADERS = {"X-Admin-Key": "test-admin-key"}
BASE = "/v1/admin/institution"


@pytest.fixture(autouse=True)
def extraction_mode(monkeypatch):
    monkeypatch.setattr(settings, "llm_provider", "extractive")


def setup(database):
    config = setup_institution(database)
    assistant = database.get(Assistant, config.assistant_id)
    assistant.active = True
    database.commit()
    return config, assistant


def test_setup_is_idempotent_and_private(client, database):
    config, assistant = setup(database)
    assert setup_institution(database).assistant_id == assistant.id
    assert database.scalar(select(func.count()).select_from(ApiKey)) == 1
    assert client.get(BASE).status_code == 422
    result = client.get(BASE, headers=HEADERS).json()
    assert result["name"] == "Assistente FAFIC"
    assert "api_key" not in str(result)
    assert "tenant_id" not in result
    assert result["groq_configured"] is False


def test_setup_refuses_ambiguous_environments(database):
    tenant = Tenant(name="FAFIC", slug="fafic")
    database.add(tenant); database.flush()
    for slug in ["sandbox", "production"]:
        database.add(Environment(tenant_id=tenant.id, name=slug, slug=slug, kind=EnvironmentKind.SANDBOX))
    database.commit()
    with pytest.raises(ValueError, match="vários ambientes"):
        setup_institution(database)
    database.rollback()
    assert database.get(InstitutionConfig, 1) is None
    assert database.scalar(select(func.count()).select_from(Environment)) == 2


def test_draft_never_displaces_published_version(client, database):
    _, assistant = setup(database)
    source = KnowledgeSource(tenant_id=assistant.tenant_id, environment_id=assistant.environment_id, assistant_id=assistant.id, name="Calendário", source_type=SourceType.DOCUMENT_URL)
    database.add(source); database.flush()
    old = KnowledgeDocument(tenant_id=assistant.tenant_id, environment_id=assistant.environment_id, assistant_id=assistant.id, source_id=source.id, title="Matrícula", content="A matrícula começa em fevereiro no portal acadêmico.", original_url="https://example.org/calendario", status=KnowledgeStatus.PUBLISHED, is_current=False)
    new = KnowledgeDocument(tenant_id=assistant.tenant_id, environment_id=assistant.environment_id, assistant_id=assistant.id, source_id=source.id, title="Matrícula", content="A matrícula começa em março no portal acadêmico.", original_url=old.original_url, version_number=2)
    database.add_all([old, new]); database.commit()
    def answer():
        return client.post(BASE + "/preview", headers=HEADERS, json={"content": "Quando começa a matrícula?"}).json()
    assert "fevereiro" in answer()["content"]
    assert client.post(f"{BASE}/documents/{new.id}/publish", headers=HEADERS).status_code == 200
    assert "março" in answer()["content"]
    assert "fevereiro" not in answer()["content"]
    assert client.post(f"{BASE}/documents/{new.id}/unpublish", headers=HEADERS).status_code == 200
    assert answer()["citations"] == []
    assert client.post(f"{BASE}/documents/{old.id}/publish", headers=HEADERS).status_code == 409


def test_widget_origin_session_binding_and_disabled_bot(client, database):
    _, assistant = setup(database)
    origin = {"Origin": settings.allowed_origins_list[0]}
    assert client.post("/v1/widget/bootstrap", json={}, headers={"Origin": "https://evil.example"}).status_code == 403
    first = client.post("/v1/widget/bootstrap", json={}, headers=origin).json()
    second = client.post("/v1/widget/bootstrap", json={}, headers=origin).json()
    headers = {**origin, "Authorization": "Bearer " + first["token"]}
    assert client.get(f"/v1/sessions/{second['session_id']}/messages", headers=headers).status_code == 404
    assert client.post("/v1/sessions", headers=headers, json={"assistant_id": str(assistant.id)}).status_code == 403
    resumed = client.post("/v1/widget/bootstrap", json={"resume_token": first["resume_token"]}, headers=origin)
    assert resumed.json()["session_id"] == first["session_id"]
    assert client.get(f"/v1/sessions/{first['session_id']}", headers={**origin, "Authorization": "Bearer " + first["resume_token"]}).status_code == 401
    assistant.active = False; database.commit()
    assert client.post(f"/v1/sessions/{first['session_id']}/messages", headers=headers, json={"content": "Olá"}).status_code == 503
    assert client.post("/v1/widget/bootstrap", json={}, headers=origin).status_code == 503
    assert client.get(BASE + "/documents", headers=HEADERS).status_code == 200
    assert client.post(BASE + "/preview", headers=HEADERS, json={"content": "Olá"}).status_code == 200


def test_fafic_does_not_hit_legacy_quota(client, database):
    _, assistant = setup(database)
    database.get(Tenant, assistant.tenant_id).monthly_message_limit = 1; database.commit()
    origin = {"Origin": settings.allowed_origins_list[0]}
    boot = client.post("/v1/widget/bootstrap", json={}, headers=origin).json()
    for _ in range(2):
        response = client.post(f"/v1/sessions/{boot['session_id']}/messages", headers={**origin, "Authorization": "Bearer " + boot["token"]}, json={"content": "Olá"})
        assert response.status_code == 201


def test_upload_validation_and_unpublished_content(client, database):
    setup(database)
    response = client.post(BASE + "/sources/upload", headers=HEADERS, files={"file": ("bad.pdf", b"invalid pdf", "application/pdf")})
    assert response.status_code == 422
    response = client.post(BASE + "/sources/upload", headers=HEADERS, files={"file": ("script.exe", b"A valid text but invalid extension", "text/plain")})
    assert response.status_code == 422
    response = client.post(BASE + "/sources/upload", headers=HEADERS, files={"file": ("horario.txt", "A biblioteca abre todos os dias às oito horas.".encode(), "text/plain")})
    assert response.status_code == 201
    answer = client.post(BASE + "/preview", headers=HEADERS, json={"content": "Qual o horário da biblioteca?"}).json()
    assert answer["citations"] == []
    documents = client.get(BASE + "/documents", headers=HEADERS).json()
    assert documents[0]["status"] == "draft"


def test_docx_tables_are_extracted():
    document = Document()
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Horário da biblioteca"
    table.cell(0, 1).text = "Segunda a sexta, das 8h às 18h"
    buffer = io.BytesIO(); document.save(buffer)
    extracted = extract_document(buffer.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "https://example.org/table.docx")
    assert "Segunda a sexta" in extracted.content


def test_institution_facade_cannot_access_other_assistant(client, database):
    setup(database)
    tenant = Tenant(name="Outra", slug="outra"); database.add(tenant); database.flush()
    environment = Environment(tenant_id=tenant.id, name="Outro", slug="outro", kind=EnvironmentKind.SANDBOX); database.add(environment); database.flush()
    assistant = Assistant(tenant_id=tenant.id, environment_id=environment.id, name="Outro", slug="outro"); database.add(assistant); database.flush()
    document = KnowledgeDocument(tenant_id=tenant.id, environment_id=environment.id, assistant_id=assistant.id, title="Privado", content="Informação privada de outra instituição.", status=KnowledgeStatus.PUBLISHED); database.add(document); database.commit()
    assert client.get(f"{BASE}/documents/{document.id}", headers=HEADERS).status_code == 404
    assert client.post(f"{BASE}/documents/{document.id}/publish", headers=HEADERS).status_code == 404
    assert client.delete(f"{BASE}/documents/{document.id}", headers=HEADERS).status_code == 404


def test_control_message_used_when_disabled_and_groq_exhausted(client, database, monkeypatch):
    import httpx
    _, assistant = setup(database)
    custom = "Nosso bot está temporariamente indisponível. Use o contato de demonstração."
    result = client.patch(BASE + "/control", headers=HEADERS, json={"unavailable_message": custom})
    assert result.status_code == 200
    assert set(result.json()) == {"active", "unavailable_message"}
    assert client.patch(BASE + "/control", headers=HEADERS, json={"unavailable_message": "  "}).status_code == 422
    origin = {"Origin": settings.allowed_origins_list[0]}
    boot = client.post("/v1/widget/bootstrap", json={}, headers=origin).json()
    client.patch(BASE + "/control", headers=HEADERS, json={"active": False})
    assert client.post("/v1/widget/bootstrap", json={}, headers=origin).json()["detail"] == custom
    assert client.get("/v1/widget/appearance", headers=origin).json()["unavailable_message"] == custom
    client.patch(BASE + "/control", headers=HEADERS, json={"active": True})
    database.add(KnowledgeDocument(tenant_id=assistant.tenant_id, environment_id=assistant.environment_id, assistant_id=assistant.id,
        title="Normativa de matrícula", content="A matrícula acontece no portal acadêmico da faculdade.", status=KnowledgeStatus.PUBLISHED))
    database.commit()
    monkeypatch.setattr(settings, "llm_provider", "groq")
    monkeypatch.setattr(settings, "groq_api_key", "fake-test-secret")
    def quota_error(*args, **kwargs):
        request = httpx.Request("POST", "https://api.groq.com")
        raise httpx.HTTPStatusError("private provider detail", request=request, response=httpx.Response(429, request=request))
    monkeypatch.setattr("app.modules.conversations.service.httpx.post", quota_error)
    response = client.post(f"/v1/sessions/{boot['session_id']}/messages", headers={**origin, "Authorization": "Bearer " + boot["token"]}, json={"content": "Como faço a matrícula?"})
    assert response.status_code == 503
    assert response.json() == {"detail": custom}
    monkeypatch.setattr(settings, "groq_api_key", "")
    assert client.post(f"/v1/sessions/{boot['session_id']}/messages", headers={**origin, "Authorization": "Bearer " + boot["token"]}, json={"content": "Como faço a matrícula?"}).json()["detail"] == custom


def test_response_is_natural_without_document_heading(client, database):
    _, assistant = setup(database)
    database.add(KnowledgeDocument(tenant_id=assistant.tenant_id, environment_id=assistant.environment_id, assistant_id=assistant.id,
        title="Documento normativo de matrícula", content="A matrícula acontece no portal acadêmico da faculdade.", status=KnowledgeStatus.PUBLISHED))
    database.commit()
    result = client.post(BASE + "/preview", headers=HEADERS, json={"content": "Como faço a matrícula?"}).json()
    assert result["content"] == "A matrícula acontece no portal acadêmico da faculdade."
    assert "Documento normativo" not in result["content"]


def test_institution_metrics_and_conversations(client, database):
    _, assistant = setup(database)
    origin = {"Origin": settings.allowed_origins_list[0]}

    # Check metrics on initial state
    res_metrics = client.get(f"{BASE}/metrics", headers=HEADERS)
    assert res_metrics.status_code == 200
    metrics_data = res_metrics.json()
    assert metrics_data["bot_name"] == assistant.name
    assert metrics_data["sessions_total"] == 0
    assert metrics_data["total_user_messages"] == 0
    assert isinstance(metrics_data["daily_messages"], list)

    # Check conversations list on initial state
    res_convs = client.get(f"{BASE}/conversations", headers=HEADERS)
    assert res_convs.status_code == 200
    assert res_convs.json()["total"] == 0
    assert res_convs.json()["conversations"] == []

    # Create a conversation and send a message
    boot = client.post("/v1/widget/bootstrap", json={}, headers=origin).json()
    session_id = boot["session_id"]
    token = boot["token"]
    msg_res = client.post(
        f"/v1/sessions/{session_id}/messages",
        headers={**origin, "Authorization": f"Bearer {token}"},
        json={"content": "Olá, bom dia!"}
    )
    assert msg_res.status_code == 201

    # Check metrics again
    updated_metrics = client.get(f"{BASE}/metrics", headers=HEADERS).json()
    assert updated_metrics["sessions_total"] == 1
    assert updated_metrics["total_user_messages"] == 1
    assert updated_metrics["confidence_average"] == 0.0

    # Check conversations list
    convs_data = client.get(f"{BASE}/conversations", headers=HEADERS).json()
    assert convs_data["total"] == 1
    assert len(convs_data["conversations"]) == 1
    c = convs_data["conversations"][0]
    assert c["id"] == session_id
    assert c["message_count"] == 2
    assert "Olá" in c["first_message"]

    # Check conversation detail
    detail = client.get(f"{BASE}/conversations/{session_id}", headers=HEADERS).json()
    assert detail["id"] == session_id
    assert len(detail["messages"]) == 2
    assert detail["messages"][0]["role"] == "user"
    assert detail["messages"][1]["role"] == "assistant"
